"""Gera o formulário de levantamento de tarefas (Excel) do projeto de automação.

Uso: python 03-propostas/projeto-automacao-ia/gerar_formulario.py
Saída: formulario-levantamento-tarefas.xlsx, na mesma pasta.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

SAIDA = Path(__file__).with_name("formulario-levantamento-tarefas.xlsx")
QTD_TAREFAS = 5

# Identidade (mesma paleta dos PDFs)
PETROLEO = "13505B"
PETROLEO_CLARO = "E3EEF0"
AREIA = "F5F0E8"
PREENCHER = "FFF3C4"  # amarelo: células para preencher
CINZA = "56676C"
FONTE = "Arial"

f = lambda **kw: Font(name=FONTE, **kw)
fill = lambda cor: PatternFill("solid", start_color=cor)
linha = Side(style="thin", color="D9E0E2")
borda = Border(left=linha, right=linha, top=linha, bottom=linha)
quebra = Alignment(wrap_text=True, vertical="top")

FREQUENCIAS = ["Diária", "Semanal", "Quinzenal", "Mensal", "Outra"]
REGRAS = ["Sempre as mesmas regras", "Às vezes exige decisão", "Exige análise quase sempre"]
SIM_NAO = ["Sim", "Não", "Não sei"]
IMPORTANCIA = ["Alta", "Média", "Baixa"]

# (chave, pergunta, ajuda, altura da linha, tipo)
# tipo: texto | numero | calc | lista:<nome>
SECOES = [
    ("1. Identificação", [
        ("nome", "Nome da tarefa", "Um nome curto. Ex.: Organizar documentos recebidos por e-mail", 30, "texto"),
        ("quem", "Quem faz hoje", "Função ou setor (não precisa de nome). Ex.: Departamento pessoal", 30, "texto"),
        ("pessoas", "Quantas pessoas fazem essa tarefa", "Número aproximado", 22, "numero"),
    ]),
    ("2. Como é feita hoje", [
        ("passos", "Passo a passo", "Descreva como explicaria para alguém novo no escritório. Pode usar frases curtas, uma por etapa.", 150, "texto"),
        ("gatilho", "O que dá início à tarefa", "Ex.: chega um e-mail, começa o mês, o cliente pede, vence um prazo", 40, "texto"),
        ("sistemas", "Sistemas e programas usados", "Ex.: sistema da folha, Excel, e-mail, site da Receita ou da prefeitura", 40, "texto"),
        ("entradas", "Documentos e arquivos usados", "Só o tipo, sem dados reais. Ex.: PDF de nota fiscal, planilha do cliente, papel", 40, "texto"),
        ("resultado", "Resultado final", "O que é entregue ao final e para quem", 40, "texto"),
    ]),
    ("3. Tempo e frequência", [
        ("freq", "Com que frequência acontece", "Escolha na lista", 22, "lista:freq"),
        ("vezes", "Quantas vezes por mês, aproximadamente", "Ex.: diária ≈ 22 · semanal ≈ 4 · mensal = 1", 22, "numero"),
        ("minutos", "Tempo gasto em cada vez (minutos)", "Uma estimativa já basta", 22, "numero"),
        ("horas", "Horas por mês gastas hoje", "Calculado automaticamente", 22, "calc"),
        ("pico", "Tem prazo ou época de pico?", "Ex.: sempre até o dia 20; aumenta no fim do ano", 40, "texto"),
    ]),
    ("4. Dificuldades", [
        ("erros", "Onde costuma dar erro ou retrabalho", "O que mais incomoda nessa tarefa", 60, "texto"),
        ("regras", "A tarefa segue regras fixas ou exige decisão?", "Escolha na lista", 22, "lista:regras"),
        ("dados", "Envolve dados pessoais? (CPF, salário, dados de funcionários de clientes)", "Escolha na lista", 30, "lista:simnao"),
    ]),
    ("5. Expectativa", [
        ("ideia", "Como imaginam essa tarefa automatizada", "Não precisa ser técnico. Ex.: os arquivos já chegarem organizados na pasta do cliente", 60, "texto"),
        ("economia", "Quantas horas por mês esperam economizar", "Uma estimativa", 22, "numero"),
        ("importancia", "Importância para o escritório", "Escolha na lista", 22, "lista:importancia"),
        ("obs", "Observações", "Qualquer detalhe que ajude a entender", 60, "texto"),
    ]),
]

EXEMPLO = {
    "nome": "Organizar documentos recebidos dos clientes por e-mail",
    "quem": "Recepção / atendimento",
    "pessoas": 2,
    "passos": "1. Abrir a caixa de e-mail do escritório.\n2. Baixar os anexos de cada cliente.\n"
              "3. Renomear no padrão: CLIENTE - TIPO - MÊS.\n4. Salvar na pasta do cliente no servidor.\n"
              "5. Avisar o responsável pelo cliente que o documento chegou.",
    "gatilho": "Chegada de e-mail de cliente com anexo",
    "sistemas": "E-mail, pastas do servidor",
    "entradas": "PDF de notas fiscais, extratos e guias; às vezes foto do celular",
    "resultado": "Documento salvo na pasta certa e responsável avisado",
    "freq": "Diária",
    "vezes": 22,
    "minutos": 30,
    "pico": "Aumenta muito nos primeiros 10 dias do mês",
    "erros": "Arquivo salvo na pasta errada ou com nome fora do padrão; e-mail esquecido",
    "regras": "Sempre as mesmas regras",
    "dados": "Sim",
    "ideia": "Os anexos já chegarem renomeados na pasta do cliente, com um aviso para o responsável",
    "economia": 10,
    "importancia": "Alta",
    "obs": "Exemplo fictício, só para mostrar como preencher.",
}


def cabecalho(ws, titulo, subtitulo):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    ws["B1"] = "Artur Vasconcelos – Tecnologia e Dados"
    ws["B1"].font = f(size=9, bold=True, color=CINZA)
    ws["B2"] = titulo
    ws["B2"].font = f(size=16, bold=True, color=PETROLEO)
    ws["B3"] = subtitulo
    ws["B3"].font = f(size=10, italic=True, color=CINZA)
    ws.row_dimensions[2].height = 26


def folha_tarefa(wb, nome_aba, titulo, subtitulo, valores=None):
    ws = wb.create_sheet(nome_aba)
    cabecalho(ws, titulo, subtitulo)
    ws.column_dimensions["B"].width = 36
    ws.column_dimensions["C"].width = 62
    ws.column_dimensions["D"].width = 44

    listas = {
        "freq": FREQUENCIAS, "regras": REGRAS, "simnao": SIM_NAO, "importancia": IMPORTANCIA,
    }
    validacoes = {}
    for chave, opcoes in listas.items():
        dv = DataValidation(type="list", formula1='"' + ",".join(opcoes) + '"', allow_blank=True,
                            showErrorMessage=True, errorTitle="Opção inválida",
                            error="Escolha uma das opções da lista.")
        ws.add_data_validation(dv)
        validacoes[chave] = dv
    num = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                         showErrorMessage=True, errorTitle="Número", error="Informe apenas um número (ex.: 4 ou 2,5).")
    ws.add_data_validation(num)

    enderecos = {}
    r = 5
    for secao, campos in SECOES:
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        c = ws.cell(r, 2, secao)
        c.font = f(size=11, bold=True, color="FFFFFF")
        c.fill = fill(PETROLEO)
        c.alignment = Alignment(vertical="center", indent=1)
        ws.row_dimensions[r].height = 22
        r += 1
        for chave, pergunta, ajuda, altura, tipo in campos:
            q = ws.cell(r, 2, pergunta)
            q.font = f(size=10, bold=True)
            q.alignment = quebra
            a = ws.cell(r, 3)
            a.alignment = quebra
            a.font = f(size=10)
            a.border = borda
            h = ws.cell(r, 4, ajuda)
            h.font = f(size=9, italic=True, color=CINZA)
            h.alignment = quebra
            ws.row_dimensions[r].height = altura
            enderecos[chave] = f"C{r}"

            if tipo == "calc":
                a.fill = fill(PETROLEO_CLARO)
                a.font = f(size=10, bold=True, color=PETROLEO)
                a.number_format = '0.0" h"'
                a.alignment = Alignment(horizontal="left", vertical="top")
            else:
                a.fill = fill(PREENCHER)
                if tipo == "numero":
                    num.add(a)
                    a.number_format = "0.##"
                    a.alignment = Alignment(horizontal="left", vertical="top")
                elif tipo.startswith("lista:"):
                    validacoes[tipo.split(":")[1]].add(a)
            r += 1
        r += 1  # respiro entre seções

    v, m = enderecos["vezes"], enderecos["minutos"]
    ws[enderecos["horas"]] = f'=IF(OR({v}="",{m}=""),"",{v}*{m}/60)'

    if valores:
        for chave, valor in valores.items():
            ws[enderecos[chave]] = valor

    ws.freeze_panes = "A5"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"B1:D{r}"
    return ws, enderecos


def folha_instrucoes(wb):
    ws = wb.active
    ws.title = "Instruções"
    cabecalho(ws, "Levantamento de tarefas para automação",
              "Preparado para Amparo Assis Contábil")
    ws.column_dimensions["B"].width = 4
    ws.column_dimensions["C"].width = 100

    blocos = [
        ("Para que serve", [
            "Este formulário é o primeiro passo do projeto de automação. Com ele, entendo como cada "
            "tarefa é feita hoje e quais têm mais potencial para serem automatizadas.",
        ]),
        ("Como preencher", [
            "Use uma aba para cada tarefa: Tarefa 1, Tarefa 2 e assim por diante.",
            "Precisa de mais abas? Clique com o botão direito em uma aba de tarefa > Mover ou copiar > marque \"Criar uma cópia\".",
            "Preencha só as células amarelas. As azuis são calculadas sozinhas.",
            "Nos campos com seta, escolha uma opção da lista.",
            "Descreva como explicaria para alguém novo no escritório. Não precisa de termos técnicos.",
            "Estimativas de tempo bastam: \"mais ou menos 30 minutos\" já ajuda muito.",
            "Não precisa preencher tudo. Um campo vazio é melhor que um chute.",
            "Veja a aba Exemplo para um modelo preenchido.",
        ]),
        ("Importante: dados de clientes", [
            "Não coloquem dados reais de clientes ou funcionários (nomes, CPF, valores). Descrevam só o tipo "
            "de documento ou informação.",
        ]),
        ("Depois de preencher", [
            "Salvem o arquivo e devolvam por e-mail para artur.augustocv@gmail.com.",
            "A aba Resumo mostra todas as tarefas juntas e o total de horas por mês.",
            "Entro em contato para conversarmos sobre as tarefas e apresentar a análise.",
        ]),
    ]

    r = 5
    for titulo, itens in blocos:
        ws.cell(r, 2, titulo).font = f(size=12, bold=True, color=PETROLEO)
        r += 1
        for item in itens:
            ws.cell(r, 2, "•").font = f(size=10, color=PETROLEO)
            c = ws.cell(r, 3, item)
            c.font = f(size=10)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[r].height = 28 if len(item) > 100 else 16
            r += 1
        r += 1

    ws.cell(r, 2, "Legenda").font = f(size=12, bold=True, color=PETROLEO)
    r += 1
    for cor, texto in [(PREENCHER, "Célula para preencher"), (PETROLEO_CLARO, "Calculado automaticamente")]:
        ws.cell(r, 2).fill = fill(cor)
        ws.cell(r, 2).border = borda
        ws.cell(r, 3, texto).font = f(size=10)
        r += 1
    r += 1

    ws.cell(r, 2, "Dúvidas").font = f(size=12, bold=True, color=PETROLEO)
    r += 1
    ws.cell(r, 3, "Artur Vasconcelos · WhatsApp (11) 9-4114-4850 · artur.augustocv@gmail.com").font = f(size=10)
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def folha_resumo(wb, tarefas):
    ws = wb.create_sheet("Resumo")
    cabecalho(ws, "Resumo das tarefas", "Preenchido automaticamente a partir das abas de tarefa")
    colunas = [
        ("Tarefa", "nome", 44, None),
        ("Frequência", "freq", 13, None),
        ("Horas/mês hoje", "horas", 15, '0.0'),
        ("Economia esperada (h/mês)", "economia", 16, '0.0'),
        ("Importância", "importancia", 13, None),
        ("Dados pessoais?", "dados", 14, None),
    ]
    for i, (titulo, _, largura, _) in enumerate(colunas):
        col = 2 + i
        c = ws.cell(5, col, titulo)
        c.font = f(size=10, bold=True, color="FFFFFF")
        c.fill = fill(PETROLEO)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[c.column_letter].width = largura
    ws.row_dimensions[5].height = 30

    r = 6
    for aba, end in tarefas:
        for i, (_, chave, _, formato) in enumerate(colunas):
            ref = f"'{aba}'!{end[chave]}"
            c = ws.cell(r, 2 + i, f'=IF({ref}="","",{ref})')
            c.font = f(size=10)
            c.border = borda
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if formato:
                c.number_format = formato
        r += 1

    ultima = r - 1
    ws.cell(r, 2, "Total").font = f(size=10, bold=True)
    for col in (4, 5):
        letra = ws.cell(r, col).column_letter
        c = ws.cell(r, col, f"=SUM({letra}6:{letra}{ultima})")
        c.font = f(size=10, bold=True, color=PETROLEO)
        c.fill = fill(PETROLEO_CLARO)
        c.number_format = '0.0" h"'
        c.border = borda
    r += 2
    nota = ws.cell(r, 2, "Se criarem novas abas de tarefa, elas não entram aqui automaticamente — "
                         "sem problema, eu considero todas na análise.")
    nota.font = f(size=9, italic=True, color=CINZA)
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def main():
    wb = Workbook()
    folha_instrucoes(wb)
    ws_ex, _ = folha_tarefa(wb, "Exemplo", "Exemplo preenchido",
                            "Tarefa fictícia, só para mostrar como preencher", EXEMPLO)
    ws_ex.sheet_properties.tabColor = "B8C4C7"
    ws_ex["C6"].comment = Comment("Exemplo fictício. Use as abas Tarefa 1, Tarefa 2… para as tarefas reais.",
                                  "Artur Vasconcelos")
    tarefas = []
    for n in range(1, QTD_TAREFAS + 1):
        aba = f"Tarefa {n}"
        ws, end = folha_tarefa(wb, aba, f"Tarefa {n}", "Preencha as células amarelas")
        ws.sheet_properties.tabColor = PETROLEO
        tarefas.append((aba, end))
    folha_resumo(wb, tarefas)
    wb.calculation.fullCalcOnLoad = True  # Excel calcula as fórmulas ao abrir
    wb.save(SAIDA)
    print("Gerado:", SAIDA)


if __name__ == "__main__":
    main()
